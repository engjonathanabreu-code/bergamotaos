"""
Importa os dados legados da Bergamota para o formato do Bergamota OS.

Fontes:
  - Planilha "Entradas e saídas 2026.xlsx"  -> vendas (cada ticket) e despesas
  - backup_confeitaria_pro_*.json            -> insumos, fichas, histórico de preços,
                                                configurações e a classificação
                                                (categoria) de cada fornecedor/despesa

Saída: scripts/seed/*.json + relatório de reconciliação contra "Saúde dos meses".

Uso:  python scripts/import_legacy.py <pasta_com_os_arquivos>
"""
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(sys.argv[1] if len(sys.argv) > 1 else '..')
OUT = Path(__file__).parent / 'seed'
OUT.mkdir(exist_ok=True)

XLSX = next(ROOT.glob('Entradas e sa*.xlsx'))
BACKUP = sorted(ROOT.glob('backup_confeitaria_pro_*.json'))[-1]
TODAY = date.today()

MONTHS = {'Jan': 1, 'Fev': 2, 'Mar': 3, 'Abr': 4, 'Mai': 5, 'Jun': 6,
          'Jul': 7, 'Ago': 8, 'Set': 9, 'Out': 10, 'Nov': 11, 'Dez': 12}
METHODS = ['dinheiro', 'pix', 'debito', 'credito']  # colunas B..E e G..J

# Grupos de despesa: (coluna data, descrição, valor, última linha de lançamento)
EXPENSE_BLOCKS = {
    'compras': ('A', 'B', 'C', 50),
    'funcionario': ('E', 'F', 'G', 16),
    'fixa': ('I', 'J', 'K', 16),
    'variavel': ('M', 'N', 'O', 16),
    'esporadica': ('Q', 'R', 'S', 16),
}
GROUP_DEFAULT_CATEGORY = {
    'compras': 'Insumos e Fornecedores',
    'funcionario': 'Funcionários',
    'fixa': 'Despesas Fixas',
    'variavel': 'Despesas Variáveis',
    'esporadica': 'Despesas Esporádicas',
}
# Na planilha de Jan–Ago o total de saídas usado na "Saúde dos meses" é
# N1 + 7000: o pró-labore dos sócios (a partir de Set aparece explícito como
# "Prolab Fran" / "Prolab Dani" – R$ 3.500 cada).
IMPLICIT_PROLABORE_MONTHS = range(1, 9)


def norm(s: str) -> str:
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


wb_v = openpyxl.load_workbook(XLSX, data_only=True)
wb_f = openpyxl.load_workbook(XLSX, data_only=False)
backup = json.load(open(BACKUP, encoding='utf-8'))

# ---------------------------------------------------------------------------
# Classificação aprendida do backup: descrição normalizada -> categoria
# ---------------------------------------------------------------------------
learned = defaultdict(Counter)
for t in backup['transactions']:
    if t['type'] == 'saida':
        learned[norm(t['description'])][t['category']] += 1


def classify(desc: str, group: str) -> str:
    n = norm(desc)
    if n in learned:
        return learned[n].most_common(1)[0][0]
    rules = [
        (r'prolab|pro labore|pro-labore', 'Pró-labore'),
        (r'simples|dare|inss|alvara|mapa|iptu|taxa', 'Impostos e Taxas'),
        (r'emprest|financ|pronampe', 'Empréstimos e Financiamentos'),
        (r'almoco|lanche|mercado', 'Alimentação'),
        (r'flor|decora|ilustra|artes|artesanato', 'Decoração e Flores'),
        (r'curso|hotmart|sebrae|nutri|tabela', 'Serviços Profissionais'),
        (r'embalag|caixa|sacola', 'Embalagens'),
        (r'celesc|luz|energia|agua|casan|unifique|internet|condominio|aluguel|seguro', 'Despesas Fixas'),
    ]
    for pat, cat in rules:
        if re.search(pat, n):
            return cat
    return GROUP_DEFAULT_CATEGORY[group]


# ---------------------------------------------------------------------------
# Vendas: cada célula numérica dentro do bloco de um dia é um ticket
# ---------------------------------------------------------------------------
sales = []          # tickets individuais
day_totals = {}     # (date) -> total conferido pela planilha (coluna L ou F)
closed_days = []
issues = []


def month_sheets(prefix):
    for ws in wb_v.worksheets:
        m = re.fullmatch(prefix + r' (\w{3}) (\d\d)', ws.title)
        if m and m.group(1) in MONTHS:
            yield ws.title, 2000 + int(m.group(2)), MONTHS[m.group(1)]


for title, year, month in month_sheets('Entradas'):
    wv, wf = wb_v[title], wb_f[title]
    starts = [r for r in range(4, wv.max_row + 1) if as_date(wv[f'A{r}'].value)]
    for i, start in enumerate(starts):
        end = (starts[i + 1] - 1) if i + 1 < len(starts) else start
        if i + 1 == len(starts):  # último bloco: vai até a linha de totais (A com fórmula)
            r = start + 1
            while r <= wv.max_row and not (isinstance(wf[f'A{r}'].value, str)):
                r += 1
            end = r - 1
        d = as_date(wv[f'A{start}'].value)
        seq = 0
        for r in range(start, end + 1):
            for col, method in zip('BCDE', METHODS):
                v = num(wv[f'{col}{r}'].value)
                if v and v > 0:
                    seq += 1
                    sales.append(dict(date=d.isoformat(), method=method, amount=round(v, 2),
                                      register='principal', seq=seq))
            if r >= start + 2:  # bloco "CZ" (mesma loja, 2º caixa)
                for col, method in zip('GHIJ', METHODS):
                    fv = wf[f'{col}{r}'].value
                    if isinstance(fv, str) and fv.startswith('='):
                        continue
                    v = num(wv[f'{col}{r}'].value)
                    if v and v > 0:
                        seq += 1
                        sales.append(dict(date=d.isoformat(), method=method, amount=round(v, 2),
                                          register='cz', seq=seq))
        ref = num(wv[f'L{start}'].value)
        if ref is None:
            ref = num(wv[f'F{start}'].value)
        day_totals[d.isoformat()] = ref or 0.0

# confere cada dia
by_day = defaultdict(float)
for s in sales:
    by_day[s['date']] += s['amount']
for d, ref in sorted(day_totals.items()):
    if abs(by_day[d] - ref) > 0.01:
        issues.append(f'Dia {d}: tickets somam {by_day[d]:.2f} mas planilha mostra {ref:.2f}')
    if ref == 0 and date.fromisoformat(d) < TODAY:
        # dia útil do calendário sem venda: provável dia fechado
        closed_days.append(dict(date=d, reason='Sem vendas registradas na planilha'))

# ---------------------------------------------------------------------------
# Despesas
# ---------------------------------------------------------------------------
transactions = []
for title, year, month in month_sheets('Saídas'):
    wv, wf = wb_v[title], wb_f[title]
    comp = f'{year}-{month:02d}'
    for group, (cd, cdesc, cval, last) in EXPENSE_BLOCKS.items():
        header = wv[f'{cdesc}2'].value
        if header is None:
            continue  # meses sem esse grupo (variável/esporádica só a partir de Set)
        for r in range(3, last + 1):
            v = num(wv[f'{cval}{r}'].value)
            desc = wv[f'{cdesc}{r}'].value
            if not v:
                continue
            d = as_date(wv[f'{cd}{r}'].value)
            note = None
            if d is None or (d.year, d.month) != (year, month):
                if d is not None:
                    note = f'Data original na planilha: {d.isoformat()}'
                d = date(year, month, 1)
            formula = wf[f'{cval}{r}'].value
            if isinstance(formula, str) and formula.startswith('='):
                note = ((note + ' · ') if note else '') + f'Composição: {formula[1:]}'
            desc = str(desc or '(sem descrição)').strip()
            transactions.append(dict(
                type='saida', date=d.isoformat(), competencia=comp,
                description=desc, category=classify(desc, group), group=group,
                value=round(v, 2), status='previsto' if d > TODAY else 'realizado',
                source='planilha', notes=note,
            ))
    if month in IMPLICIT_PROLABORE_MONTHS:
        for socio in ('Fran', 'Dani'):
            transactions.append(dict(
                type='saida', date=f'{comp}-01', competencia=comp,
                description=f'Prolab {socio}', category='Pró-labore', group='funcionario',
                value=3500.0, status='realizado', source='planilha',
                notes='Pró-labore implícito (+7.000 no total de saídas da planilha)',
            ))

# ---------------------------------------------------------------------------
# Merge com o backup: só entra o que a planilha não tem
# ---------------------------------------------------------------------------
pool = Counter((t['competencia'], round(t['value'], 2)) for t in transactions)
# a planilha às vezes lança numa aba datas de outro mês; casa também por descrição+valor
pool_desc = Counter((norm(t['description']), round(t['value'], 2)) for t in transactions)
added_from_backup = 0
for t in backup['transactions']:
    if t['type'] == 'entrada':
        d = t['date']
        if by_day.get(d, 0) > 0 or t.get('isHoliday'):
            if t.get('isHoliday'):
                closed_days.append(dict(date=d, reason=t.get('holidayReason') or 'Feriado'))
            continue
        # venda do dia que não está na planilha: entra como lançamento consolidado
        bd = t.get('paymentBreakdown') or {'outro': t['value']}
        for m, v in bd.items():
            if v:
                sales.append(dict(date=d, method=m, amount=round(v, 2), register='principal',
                                  seq=0, aggregate=True))
        added_from_backup += 1
        continue
    key = (t['date'][:7], round(t['value'], 2))
    key2 = (norm(t['description']), round(t['value'], 2))
    if pool[key] > 0:
        pool[key] -= 1
        continue
    if pool_desc[key2] > 0:
        pool_desc[key2] -= 1
        continue
    d = date.fromisoformat(t['date'])
    transactions.append(dict(
        type='saida', date=t['date'], competencia=t['date'][:7], description=t['description'],
        category=t['category'], group=t.get('group') or 'compras', value=round(t['value'], 2),
        status='previsto' if d > TODAY else 'realizado', source='backup',
        notes='Veio só do backup JSON (não encontrado na planilha) — conferir',
    ))
    added_from_backup += 1

# ---------------------------------------------------------------------------
# Reconciliação com "Saúde dos meses"
# ---------------------------------------------------------------------------
health = wb_v['Saúde dos meses']
names = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho', 'agosto',
         'setembro', 'outubro', 'novembro', 'dezembro']
inc = defaultdict(float)
for s in sales:
    inc[s['date'][:7]] += s['amount']
exp = defaultdict(float)
for t in transactions:
    if t['source'] == 'planilha':
        exp[t['competencia']] += t['value']
report = []
for r in range(2, 14):
    mname = norm(health[f'A{r}'].value)
    m = names.index(mname) + 1
    comp = f'2026-{m:02d}'
    b, c = num(health[f'B{r}'].value), num(health[f'C{r}'].value)
    line = f'{comp}: entradas {inc[comp]:>10.2f} (planilha {b or 0:>10.2f})  saídas {exp[comp]:>10.2f} (planilha {c or 0:>10.2f})'
    ok = (b is None or abs(inc[comp] - b) < 0.05) and (c is None or abs(exp[comp] - c) < 0.05)
    report.append(('OK  ' if ok else 'DIF ') + line)

# ---------------------------------------------------------------------------
# Cadastros vindos do backup
# ---------------------------------------------------------------------------
settings = backup['settings']
for extra in ['Pró-labore', 'Despesas Variáveis', 'Despesas Esporádicas']:
    if extra not in settings['expenseCategories']:
        settings['expenseCategories'].append(extra)

dump = lambda name, obj: json.dump(obj, open(OUT / f'{name}.json', 'w', encoding='utf-8'),
                                   ensure_ascii=False, indent=1)
dump('ingredients', backup['ingredients'])
dump('recipes', backup['recipes'])
dump('price_history', backup['priceHistory'])
dump('settings', settings)
dump('sales', sales)
dump('transactions', transactions)
seen = set()
dump('closed_days', [c for c in closed_days if not (c['date'] in seen or seen.add(c['date']))])

print(f'Tickets de venda: {len(sales)}  | Despesas: {len(transactions)}  | Do backup: {added_from_backup}')
print(f'Insumos: {len(backup["ingredients"])}  Fichas: {len(backup["recipes"])}  Hist. preços: {len(backup["priceHistory"])}')
print('\n'.join(report))
print(f'Dias fechados/sem venda: {len(seen)}')
if issues:
    print('Divergências por dia:')
    print('\n'.join(issues[:30]))
print('Categorias:', Counter(t['category'] for t in transactions).most_common())
